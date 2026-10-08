"""
Refunds handled inside Withdrawals — Task 11
==============================================
The dedicated Refunds page was removed (Task 9) — oversubscription refunds
created by the Shares deployment allocator (Task 3) must show up in, and be
processable from, the plain Withdrawals screen: GET /api/v1/withdrawals (with
its new `product` filter), the bulk approve/reject scope, and the Excel export.

Run with: pytest tests/test_withdrawals_refunds.py -v
"""

import json
from datetime import datetime, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from io import BytesIO

from app.database.database import db
from app.Batch.core_fund import CoreFund
from app.Funds.share_class import ShareClass
from app.Batch.model import Batch
from app.Investments.model import Investment, Withdrawal, WITHDRAWAL_SOURCE_OVERSUBSCRIPTION_REFUND
from app.logic.batch_management_service import BatchManagementService
from app.logic.shares_transfer_service import SharesTransferService
from app.logic.shares_deployment_service import SharesDeploymentService


def get_auth_headers(token):
    return {"Authorization": f"Bearer {token}"}


def _make_fund_and_class(fund_name, fund_code, class_code, currency):
    fund = CoreFund(fund_name=fund_name, fund_code=fund_code)
    db.session.add(fund)
    db.session.flush()
    share_class = ShareClass(core_fund_id=fund.id, class_name=f"{fund_name} Class", class_code=class_code, currency=currency)
    db.session.add(share_class)
    db.session.flush()
    return fund, share_class


def _investment(batch, fund, code, amount, name=None):
    """Mirrors test_shares_deployment.py's own `_investment` helper, but also
    sets fund_id — every real investment-creation path does (see Withdrawal
    rows elsewhere relying on it), so an accurate refund-product-filter test
    needs it set here too."""
    inv = Investment(
        batch_id=batch.id,
        fund_id=fund.id,
        investor_name=name or code,
        investor_email=f"{code.lower()}@example.com",
        internal_client_code=code,
        amount_deposited=Decimal(amount),
        share_class_id=batch.share_class_id,
    )
    db.session.add(inv)
    db.session.commit()
    return inv


def _make_oversubscribed_dprp_refund(app):
    """Worked example from test_shares_deployment.py's
    test_oversubscribed_both_rounds: A=2000, B=2000, C=6000 USD; entry 2%, tx
    cost 98; deployed at rate 1500, price 525, 20000 shares allocated ->
    investor C gets refunded $2702.00 (oversubscribed)."""
    dprp_fund, share_class = _make_fund_and_class(
        "Dangote Petroleum Refinery and Petrochemicals FZE", "DPRP-WD", "DPRP-WD", "USD"
    )
    batch = BatchManagementService.find_or_create_batch(share_class.id)
    a = _investment(batch, dprp_fund, "REFUND-A", "2000.00")
    b = _investment(batch, dprp_fund, "REFUND-B", "2000.00")
    c = _investment(batch, dprp_fund, "REFUND-C", "6000.00")
    SharesTransferService.apply_to_batch(batch, db.session, entry_fee_percent="2", total_transaction_cost_usd="98")
    db.session.expire_all()
    batch = db.session.get(Batch, batch.id)
    SharesDeploymentService.apply_to_batch(batch, db.session, "1500", "525", "20000")
    db.session.commit()
    return dprp_fund, batch, a, b, c


def _make_funds_investor_withdrawal(app):
    """A plain FUNDS investor-request withdrawal, for cross-product isolation checks."""
    fund, share_class = _make_fund_and_class("Axiom", "AXIOM-WD", "AXIOM-WD", "USD")
    batch = BatchManagementService.find_or_create_batch(share_class.id)
    inv = _investment(batch, fund, "FUNDS-INV", "5000.00")
    w = Withdrawal(
        internal_client_code=inv.internal_client_code,
        fund_id=fund.id,
        fund_name=fund.fund_name,
        gross_amount=Decimal("500.00"),
        amount=Decimal("500.00"),
        date_withdrawn=datetime.now(timezone.utc),
        status="Pending",
    )
    db.session.add(w)
    db.session.commit()
    return fund, w


@pytest.mark.unit
class TestRefundsAppearInWithdrawals:
    def test_shares_deployment_refund_appears_with_oversubscription_refund_type(self, client, auth_token, app):
        with app.app_context():
            dprp_fund, batch, a, b, c = _make_oversubscribed_dprp_refund(app)

            resp = client.get("/api/v1/withdrawals", headers=get_auth_headers(auth_token))
            assert resp.status_code == 200
            rows = resp.get_json()["data"]
            refund_rows = [r for r in rows if r["source"] == WITHDRAWAL_SOURCE_OVERSUBSCRIPTION_REFUND]
            assert len(refund_rows) == 1

            row = refund_rows[0]
            assert row["client_id"] == "REFUND-C"
            assert row["investor_name"] == "REFUND-C"
            assert row["amount"] == pytest.approx(2702.00)
            assert row["status"] == "Pending"
            assert row["product_type"] == "SHARES"
            assert row["batch_name"] == batch.batch_name
            assert row["rate_used"] == pytest.approx(1500.0)
            assert row["date_withdrawn"] is not None

    def test_no_refund_for_fully_allocated_investors(self, app, client, auth_token):
        with app.app_context():
            _, batch, a, b, c = _make_oversubscribed_dprp_refund(app)
            resp = client.get(
                f"/api/v1/withdrawals?search=REFUND-A", headers=get_auth_headers(auth_token)
            )
            rows = resp.get_json()["data"]
            assert all(r["amount"] == 0 or r["client_id"] != "REFUND-A" for r in rows) or len(rows) == 0


@pytest.mark.unit
class TestWithdrawalFilters:
    def test_product_filter_shares_returns_only_refund(self, client, auth_token, app):
        with app.app_context():
            _make_oversubscribed_dprp_refund(app)
            _make_funds_investor_withdrawal(app)

            resp = client.get("/api/v1/withdrawals?product=SHARES", headers=get_auth_headers(auth_token))
            assert resp.status_code == 200
            rows = resp.get_json()["data"]
            assert len(rows) == 1
            assert rows[0]["product_type"] == "SHARES"
            assert rows[0]["source"] == WITHDRAWAL_SOURCE_OVERSUBSCRIPTION_REFUND

    def test_product_filter_funds_excludes_refund(self, client, auth_token, app):
        with app.app_context():
            _make_oversubscribed_dprp_refund(app)
            _make_funds_investor_withdrawal(app)

            resp = client.get("/api/v1/withdrawals?product=FUNDS", headers=get_auth_headers(auth_token))
            assert resp.status_code == 200
            rows = resp.get_json()["data"]
            assert len(rows) == 1
            assert rows[0]["product_type"] == "FUNDS"
            assert rows[0]["source"] != WITHDRAWAL_SOURCE_OVERSUBSCRIPTION_REFUND

    def test_source_filter_still_isolates_refund_type(self, client, auth_token, app):
        with app.app_context():
            _make_oversubscribed_dprp_refund(app)
            _make_funds_investor_withdrawal(app)

            resp = client.get(
                "/api/v1/withdrawals?source=oversubscription_refund", headers=get_auth_headers(auth_token)
            )
            rows = resp.get_json()["data"]
            assert len(rows) == 1
            assert rows[0]["source"] == WITHDRAWAL_SOURCE_OVERSUBSCRIPTION_REFUND

    def test_combined_product_and_source_filter(self, client, auth_token, app):
        with app.app_context():
            _make_oversubscribed_dprp_refund(app)
            _make_funds_investor_withdrawal(app)

            resp = client.get(
                "/api/v1/withdrawals?product=FUNDS&source=oversubscription_refund",
                headers=get_auth_headers(auth_token),
            )
            rows = resp.get_json()["data"]
            assert rows == []


@pytest.mark.unit
class TestApprovalFlowUnchanged:
    """Task 11: 'The same maker-checker rules apply ... No fees are refundable.
    Refunds are paid in USD.' — verifies the existing single and bulk
    approve/reject pipeline still works unmodified for a refund row."""

    def test_single_approve_sets_status_and_timestamp_no_fee_deducted(self, client, auth_token, app):
        with app.app_context():
            _make_oversubscribed_dprp_refund(app)
            refund = db.session.query(Withdrawal).filter(
                Withdrawal.source == WITHDRAWAL_SOURCE_OVERSUBSCRIPTION_REFUND
            ).first()
            refund_id = refund.id
            assert refund.exit_fee_deducted == Decimal("0.00")
            assert refund.gross_amount == refund.amount  # no fee taken

            resp = client.patch(
                f"/api/v1/withdrawals/{refund_id}",
                data=json.dumps({"status": "Approved"}),
                headers={**get_auth_headers(auth_token), "Content-Type": "application/json"},
            )
            assert resp.status_code == 200

            updated = db.session.get(Withdrawal, refund_id)
            assert updated.status == "Approved"
            assert updated.approved_at is not None

    def test_bulk_action_scoped_to_shares_only_touches_refund(self, client, auth_token, app):
        with app.app_context():
            _make_oversubscribed_dprp_refund(app)
            _, funds_withdrawal = _make_funds_investor_withdrawal(app)
            funds_withdrawal_id = funds_withdrawal.id

            resp = client.post(
                "/api/v1/withdrawals/bulk-action",
                data=json.dumps({"action": "approve", "scope": {"product": "SHARES"}}),
                headers={**get_auth_headers(auth_token), "Content-Type": "application/json"},
            )
            assert resp.status_code == 200
            processed = resp.get_json()["data"]["processed"]
            assert processed == 1

            refund = db.session.query(Withdrawal).filter(
                Withdrawal.source == WITHDRAWAL_SOURCE_OVERSUBSCRIPTION_REFUND
            ).first()
            assert refund.status == "Approved"

            still_pending = db.session.get(Withdrawal, funds_withdrawal_id)
            assert still_pending.status == "Pending"


@pytest.mark.unit
class TestWithdrawalsExcelExportIncludesProductAndType:
    def test_excel_export_has_product_and_source_columns(self, client, auth_token, app):
        with app.app_context():
            _make_oversubscribed_dprp_refund(app)
            _make_funds_investor_withdrawal(app)

            resp = client.get("/api/v1/reports/withdrawals/excel", headers=get_auth_headers(auth_token))
            assert resp.status_code == 200

            wb = load_workbook(BytesIO(resp.data))
            ws = wb["Withdrawals"]
            headers = [cell.value for cell in ws[1]]
            assert "product" in headers
            assert "source" in headers

            product_col = headers.index("product") + 1
            source_col = headers.index("source") + 1
            values = {
                (ws.cell(row=r, column=source_col).value): ws.cell(row=r, column=product_col).value
                for r in range(2, ws.max_row + 1)
            }
            assert values.get("Oversubscription Refund") == "Shares"
            assert values.get("Investor Request") == "Funds"
