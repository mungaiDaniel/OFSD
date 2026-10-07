"""
"Show all" + server-side Excel export — Task 6
================================================
Investor directory and batch client lists stay paginated/fast by default,
but support a "show all" mode (chunked/paginated fetching the frontend
virtualizes, never one giant fetch) and a server-generated "Export to
Excel" that covers every row for the current scope/filters, not just the
visible page.

Covers: a large seeded dataset stays responsive under "show all" paging,
and each export's row count matches the on-screen total. Also checks the
export is audit-logged (who/when/scope/row count).

Run with: pytest tests/test_show_all_and_export.py -v
"""

import time
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest

from app.database.database import db
from app.database.product_type import FUNDS, SHARES
from app.Batch.core_fund import CoreFund
from app.Batch.model import Batch
from app.Investments.model import Investment
from app.utils.audit_log import AuditLog

LARGE_N = 3000


def get_auth_headers(token):
    return {'Authorization': f'Bearer {token}'}


def _data_row_count(xlsx_bytes):
    """openpyxl write-only workbooks don't cache worksheet dimensions, so
    ws.max_row comes back None when read back with read_only=True — count by
    iterating instead."""
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True)
    ws = wb.active
    return sum(1 for _ in ws.iter_rows(min_row=2))


def _bulk_seed_investors(fund_id, batch_id, n, code_prefix, product_type):
    """Fast bulk insert — bypasses the ORM (and its mapper events), so
    product_type is set explicitly per row rather than inherited from batch."""
    now = datetime.now(timezone.utc)
    rows = [
        {
            "investor_name": f"Investor {i}",
            "investor_email": f"investor{i}@example.com",
            "internal_client_code": f"{code_prefix}-{i:06d}",
            "amount_deposited": Decimal("1000.00"),
            "deployment_fee_deducted": Decimal("0.00"),
            "transfer_fee_deducted": Decimal("0.00"),
            "oversubscription_refund_amount": Decimal("0.00"),
            "date_deposited": now,
            "batch_id": batch_id,
            "fund_id": fund_id,
            "product_type": product_type,
        }
        for i in range(n)
    ]
    db.session.bulk_insert_mappings(Investment, rows)
    db.session.commit()


@pytest.mark.unit
class TestShowAllStaysResponsive:

    def test_investor_directory_show_all_chunk_is_fast_at_scale(self, client, auth_token, app):
        with app.app_context():
            fund = CoreFund(fund_name="Atium")
            db.session.add(fund)
            db.session.flush()
            batch = Batch(batch_name="Atium Batch", date_deployed=datetime(2026, 1, 1, tzinfo=timezone.utc))
            db.session.add(batch)
            db.session.flush()
            _bulk_seed_investors(fund.id, batch.id, LARGE_N, "DIR", FUNDS)

        # First "show all" chunk — the frontend loads 200 at a time and virtualizes.
        started = time.monotonic()
        response = client.get(
            "/api/v1/investors?page=1&per_page=200",
            headers=get_auth_headers(auth_token),
        )
        elapsed = time.monotonic() - started

        assert response.status_code == 200, response.get_data(as_text=True)
        data = response.get_json()
        assert data["pagination"]["total"] == LARGE_N
        assert len(data["data"]) == 200
        # Not a hard SLA — just proof a single page's cost doesn't scale with
        # the full 3,000-row table (the old failure mode this replaces).
        assert elapsed < 10, f"single page took {elapsed:.2f}s against a {LARGE_N}-row table"

    def test_batch_show_all_chunk_is_fast_at_scale(self, client, auth_token, app):
        with app.app_context():
            fund = CoreFund(fund_name="Atium")
            db.session.add(fund)
            db.session.flush()
            batch = Batch(batch_name="Atium Batch", date_deployed=datetime(2026, 1, 1, tzinfo=timezone.utc))
            db.session.add(batch)
            db.session.flush()
            _bulk_seed_investors(fund.id, batch.id, LARGE_N, "BATCH", FUNDS)
            # Simulate the normal warm-cache state (refreshed at write time by every
            # deploy/withdrawal/valuation event in production — see
            # BatchController.refresh_cached_total_capital) rather than exercising
            # its separate, pre-existing "cold cache" self-heal fallback, which
            # does its own full-batch recompute and isn't what Task 6 is about.
            batch.cached_total_capital = Decimal(str(LARGE_N * 1000))
            db.session.commit()
            batch_id = batch.id

        started = time.monotonic()
        response = client.get(
            f"/api/v1/batches/{batch_id}?show_all=1&page=1&per_page=200",
            headers=get_auth_headers(auth_token),
        )
        elapsed = time.monotonic() - started

        assert response.status_code == 200, response.get_data(as_text=True)
        data = response.get_json()["data"]
        assert data["investments_pagination"]["total"] == LARGE_N
        assert len(data["investments"]) == 200
        assert elapsed < 10, f"single page took {elapsed:.2f}s against a {LARGE_N}-row batch"

        # Without show_all (and no search), the pre-existing fast/empty behavior
        # for the batch header view is unchanged.
        response2 = client.get(f"/api/v1/batches/{batch_id}", headers=get_auth_headers(auth_token))
        data2 = response2.get_json()["data"]
        assert data2["investments"] == []
        assert data2["investors_count"] == LARGE_N


@pytest.mark.unit
class TestExcelExportMatchesOnScreenTotal:

    def test_investor_directory_export_row_count_matches_total(self, client, auth_token, app):
        with app.app_context():
            fund = CoreFund(fund_name="Atium")
            db.session.add(fund)
            db.session.flush()
            batch = Batch(batch_name="Atium Batch", date_deployed=datetime(2026, 1, 1, tzinfo=timezone.utc))
            db.session.add(batch)
            db.session.flush()
            _bulk_seed_investors(fund.id, batch.id, LARGE_N, "EXPDIR", FUNDS)

        list_resp = client.get("/api/v1/investors?page=1&per_page=1", headers=get_auth_headers(auth_token))
        on_screen_total = list_resp.get_json()["pagination"]["total"]
        assert on_screen_total == LARGE_N

        export_resp = client.get("/api/v1/investors/export", headers=get_auth_headers(auth_token))
        assert export_resp.status_code == 200
        assert "spreadsheetml" in export_resp.content_type

        assert _data_row_count(export_resp.data) == on_screen_total

    def test_batch_investments_export_row_count_matches_total(self, client, auth_token, app):
        with app.app_context():
            fund = CoreFund(fund_name="Dangote Petroleum Refinery and Petrochemicals FZE", fund_code="DPRP")
            db.session.add(fund)
            db.session.flush()
            batch = Batch(batch_name="DPRP Batch", product_type=SHARES, date_deployed=datetime(2026, 1, 1, tzinfo=timezone.utc))
            db.session.add(batch)
            db.session.flush()
            _bulk_seed_investors(fund.id, batch.id, LARGE_N, "EXPBATCH", SHARES)
            batch.cached_total_capital = Decimal(str(LARGE_N * 1000))
            db.session.commit()
            batch_id = batch.id

        batch_resp = client.get(f"/api/v1/batches/{batch_id}?show_all=1&per_page=1", headers=get_auth_headers(auth_token))
        on_screen_total = batch_resp.get_json()["data"]["investments_pagination"]["total"]
        assert on_screen_total == LARGE_N

        export_resp = client.get(f"/api/v1/batches/{batch_id}/investments/export", headers=get_auth_headers(auth_token))
        assert export_resp.status_code == 200

        assert _data_row_count(export_resp.data) == on_screen_total

        wb = openpyxl.load_workbook(BytesIO(export_resp.data), read_only=True)
        ws = wb.active
        header = next(iter(ws.iter_rows(values_only=True)))
        assert "Net (NGN)" in header
        assert "Shares Allocated" in header
        assert "Refund (USD)" in header

    def test_export_respects_search_filter_not_just_total(self, client, auth_token, app):
        with app.app_context():
            fund = CoreFund(fund_name="Atium")
            db.session.add(fund)
            db.session.flush()
            batch = Batch(batch_name="Atium Batch", date_deployed=datetime(2026, 1, 1, tzinfo=timezone.utc))
            db.session.add(batch)
            db.session.flush()
            _bulk_seed_investors(fund.id, batch.id, 50, "FILT", FUNDS)

        filtered_list = client.get(
            "/api/v1/investors?page=1&per_page=1&search=FILT-000001",
            headers=get_auth_headers(auth_token),
        )
        assert filtered_list.get_json()["pagination"]["total"] == 1

        filtered_export = client.get(
            "/api/v1/investors/export?search=FILT-000001",
            headers=get_auth_headers(auth_token),
        )
        assert _data_row_count(filtered_export.data) == 1

    def test_export_is_audit_logged_with_scope_and_row_count(self, client, auth_token, app):
        with app.app_context():
            fund = CoreFund(fund_name="Atium")
            db.session.add(fund)
            db.session.flush()
            batch = Batch(batch_name="Atium Batch", date_deployed=datetime(2026, 1, 1, tzinfo=timezone.utc))
            db.session.add(batch)
            db.session.flush()
            _bulk_seed_investors(fund.id, batch.id, 10, "AUDIT", FUNDS)
            batch_id = batch.id

        client.get("/api/v1/investors/export?product=FUNDS", headers=get_auth_headers(auth_token))
        client.get(f"/api/v1/batches/{batch_id}/investments/export", headers=get_auth_headers(auth_token))

        with app.app_context():
            dir_log = (
                db.session.query(AuditLog)
                .filter(AuditLog.action == "EXCEL_EXPORT", AuditLog.target_type == "investor_directory")
                .order_by(AuditLog.id.desc())
                .first()
            )
            assert dir_log is not None
            assert dir_log.user_email == "admin@example.com"
            assert dir_log.timestamp is not None
            assert '"row_count": 10' in dir_log.new_value
            assert '"scope": "FUNDS"' in dir_log.new_value

            batch_log = (
                db.session.query(AuditLog)
                .filter(AuditLog.action == "EXCEL_EXPORT", AuditLog.target_type == "batch", AuditLog.target_id == batch_id)
                .order_by(AuditLog.id.desc())
                .first()
            )
            assert batch_log is not None
            assert batch_log.user_email == "admin@example.com"
            assert '"row_count": 10' in batch_log.new_value
