from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook

from app.database.database import db
from app.Investments.model import Investment
from app.Batch.core_fund import CoreFund


class ExcelUploadHandler:
    """
    Compatibility utility for legacy fund upload controller imports.

    This module restores the expected interface used by:
    - app.Batch.fund_controllers.ExcelUploadController
    """

    HEADER_ALIASES = {
        "investor_name": {"investor name", "client name", "name"},
        "investor_email": {"email", "investor email"},
        "investor_phone": {"phone", "phone number", "investor phone"},
        "internal_client_code": {"internal client code", "internal_client_code", "client code", "code"},
        "amount_deposited": {"amount deposited", "amount(usd)", "amount", "deposit amount"},
        "date_deposited": {"date deposited", "deposit date", "date_deposited"},
        "wealth_manager": {"wealth manager"},
        "IFA": {"ifa", "independent financial advisor"},
        "contract_note": {"contract note", "contract url"},
        "fund_name": {"fund", "fund name", "fund_name"},
    }

    @classmethod
    def parse_excel_file(cls, file_path: str) -> Tuple[bool, List[Dict[str, Any]], str]:
        try:
            wb = load_workbook(file_path, data_only=True)
            ws = wb.active
        except Exception as exc:
            return False, [], f"Failed to read Excel file: {exc}"

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return False, [], "Excel file is empty"

        raw_headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        mapped_headers: Dict[int, str] = {}
        for idx, header in enumerate(raw_headers):
            h = header.lower().strip()
            for canonical, aliases in cls.HEADER_ALIASES.items():
                if h in aliases:
                    mapped_headers[idx] = canonical
                    break

        required = {"investor_name", "internal_client_code", "amount_deposited"}
        if not required.issubset(set(mapped_headers.values())):
            return False, [], "Missing required columns: Investor Name, Internal Client Code, Amount Deposited"

        data: List[Dict[str, Any]] = []
        for r in rows[1:]:
            if r is None:
                continue
            row_data: Dict[str, Any] = {}
            for idx, value in enumerate(r):
                key = mapped_headers.get(idx)
                if key:
                    row_data[key] = value

            # Skip fully blank rows
            if not any(v is not None and str(v).strip() != "" for v in row_data.values()):
                continue

            data.append(row_data)

        if not data:
            return False, [], "No valid data rows found in Excel file"

        return True, data, "Excel parsed successfully"

    @classmethod
    def auto_assign_funds(cls, excel_data: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        for row in excel_data:
            fund_name = (str(row.get("fund_name") or "")).strip()
            if not fund_name:
                row["fund_name"] = "Default"
        return excel_data

    @classmethod
    def bulk_upload_investments(
        cls,
        batch_id: int,
        excel_data: List[Dict[str, Any]],
        _batch_date_deployed: Any = None,
    ) -> Tuple[bool, int, str, Dict[str, Dict[str, Any]]]:
        try:
            created_count = 0
            fund_summary: Dict[str, Dict[str, Any]] = {}

            for row in excel_data:
                investor_name = str(row.get("investor_name") or "").strip()
                internal_code = str(row.get("internal_client_code") or "").strip()
                if not investor_name or not internal_code:
                    continue

                amount_raw = row.get("amount_deposited")
                try:
                    amount = Decimal(str(amount_raw if amount_raw is not None else "0")).quantize(Decimal("0.01"))
                except (InvalidOperation, ValueError):
                    amount = Decimal("0.00")
                if amount <= 0:
                    continue

                fund_name = str(row.get("fund_name") or "Default").strip() or "Default"
                fund = db.session.query(CoreFund).filter(CoreFund.fund_name == fund_name).first()
                if not fund:
                    fund = CoreFund(fund_name=fund_name, fund_code=None, is_active=True)
                    db.session.add(fund)
                    db.session.flush()

                email = str(row.get("investor_email") or "").strip() or f"{internal_code.lower()}@placeholder.local"
                phone = str(row.get("investor_phone") or "").strip() or None

                date_deposited = row.get("date_deposited")
                if isinstance(date_deposited, datetime):
                    date_deposited_val = date_deposited
                else:
                    date_deposited_val = datetime.now(timezone.utc)

                inv = Investment(
                    investor_name=investor_name,
                    investor_email=email,
                    investor_phone=phone,
                    internal_client_code=internal_code,
                    amount_deposited=amount,
                    date_deposited=date_deposited_val,
                    batch_id=batch_id,
                    fund_id=fund.id,
                    fund_name=fund_name,
                    wealth_manager=(str(row.get("wealth_manager") or "").strip() or None),
                    IFA=(str(row.get("IFA") or "").strip() or None),
                    contract_note=(str(row.get("contract_note") or "").strip() or None),
                )
                db.session.add(inv)
                created_count += 1

                if fund_name not in fund_summary:
                    fund_summary[fund_name] = {
                        "fund_id": fund.id,
                        "investor_count": 0,
                        "total_capital": 0.0,
                    }
                fund_summary[fund_name]["investor_count"] += 1
                fund_summary[fund_name]["total_capital"] += float(amount)

            if created_count == 0:
                db.session.rollback()
                return False, 0, "No valid investment rows found in uploaded file", {}

            db.session.commit()
            return True, created_count, f"Uploaded {created_count} investment(s) successfully", fund_summary
        except Exception as exc:
            db.session.rollback()
            return False, 0, f"Bulk upload failed: {exc}", {}
